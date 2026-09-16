# -*- coding: utf-8 -*-
"""Đọc 2 bảng CHẤM CÔNG trên Drive -> % Thưởng BC02.

Bảng 1 "CHẤM CÔNG NT/TK" (file Excel .xlsx "CHẤM CÔNG NT/TK_2026_GreenVita",
GOOGLE_SHEET_ID_CHAM_CONG): mỗi tháng 1 tab BCC.5, BCC.6 ... BCC.9.
Bảng 2 "Chấm công OCP" (Google Sheet "Chấm công OCP 2026", GOOGLE_SHEET_ID_CHAM_CONG_2,
xuất ra .xlsx để đọc cùng 1 cách): tab "BCC T5" ... "BCC T7", "Tháng 8", "Tháng 9".
Nhân viên KHÔNG có tên ở bảng 1 thì tìm sang bảng 2, CÙNG 1 CÁCH TÍNH. Khối "CHẤM CÔNG"
dưới bảng BC02 có cột "Bảng chấm công" ghi tên bảng đã lấy (config.CHAM_CONG_TEN_BANG).
Bảng nào không đọc được (chưa chia sẻ, chưa có tab tháng ...) chỉ cảnh báo, vẫn dùng bảng kia.

Cấu trúc tab (2 bảng giống nhau): dòng header có "HỌ VÀ TÊN" và ô "THÁNG <số>" /
"NĂM <số>"; dòng ngay dưới là ngày của từng cột; nhân viên bắt đầu từ dòng có STT là số
(tab có thể có thêm khối phụ phía dưới - vd "NHÂN SỰ TÂM AN" - STT đánh lại từ 1, vẫn đọc
như thường; tên trùng thì lấy dòng ĐẦU). Tab của tháng chọn theo TÊN TAB ("BCC.9" /
"BCC T9" / "Tháng 9") và ô THÁNG/NĂM; hai cái lệch nhau (tab "Tháng 9" của OCP quên sửa
ô THÁNG nên dòng ngày vẫn là tháng 8) thì tin theo TÊN TAB: các cột hiểu là ngày 1..31
của tháng đó và ghi cảnh báo. Ô ngày ghi:
    8 / 7.5 / 4 ...  số GIỜ làm trong ngày (đủ công = 8 giờ)
    KL   nghỉ không lương            P / P/2  nghỉ phép (có lương)
    NL   nghỉ lễ hưởng lương          CĐ       nghỉ chế độ hưởng lương
    X / M đủ công (đi muộn đủ công)   X/2, M/2 nửa công
    HV   học việc (tính như đi làm)   CN, CN/2 đi làm Chủ nhật (cả / nửa ngày)
    trống = không phải ngày làm (Chủ nhật, chưa vào làm, ngày chưa tới)

QUY TẮC % THƯỞNG (config.CHAM_CONG_*):
    - Ngày nghỉ KHÔNG LƯƠNG (KL = 1 ngày; KL/2, X/2, M/2 = nửa ngày) được cộng dồn.
    - Ngày làm KHÔNG ĐỦ 8 giờ: số giờ thiếu cộng dồn, đủ 8 giờ = 1 ngày nghỉ không lương.
    - Cứ đủ 2 ngày nghỉ (quy đổi) -> trừ 10% thưởng; 4 ngày -> 20% ... (không âm).
    - Nghỉ phép (P), nghỉ lễ (NL), nghỉ chế độ (CĐ) có lương -> KHÔNG trừ.
Tên khớp giữa Pancake và chấm công theo tiền tố, bỏ dấu (cùng quy tắc với Lịch trực);
khớp được ở cả 2 bảng thì lấy bảng 1, trừ khi bảng 2 khớp tên DÀI hơn (đúng người hơn).
Không khớp ai ở cả 2 bảng (đã nghỉ, tên viết khác ...) -> giữ % cũ trên sheet / 100%.
"""
from __future__ import annotations

import io
import math
import re
import time
import unicodedata
from datetime import date

from google.auth.transport.requests import AuthorizedSession
from google.oauth2.service_account import Credentials

import config
from lich_truc import chuan_hoa_ten, khop_ten

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_GSHEET_MIME = "application/vnd.google-apps.spreadsheet"
_NGAY_RE = re.compile(r"^\s*(\d{1,2})\s*/\s*(\d{1,2})")
# Tên tab sau _ma(): "BCC.9" / "BCCT9" / "THANG9" -> tháng 9
_TEN_TAB_RE = re.compile(r"^(?:BCC\.?T?|THANG)0?(\d{1,2})$")
_RETRY_STATUS = {429, 500, 502, 503}
_RETRY_DELAYS = (15, 30, 60)

# Thứ tự tìm tên: bảng 1 (NT/TK) trước, không thấy mới sang bảng 2 (OCP)
CAC_BANG = (1, 2)

# Khối "CHẤM CÔNG" ghi dưới bảng BC02 (build_block): header + các cột xuống dòng khi ghi sheet
KHOI_HEADER = ["STT", "Tên trên bảng thưởng", "Tên trên chấm công", "Bảng chấm công",
               "Bộ phận (chấm công)", "Ngày nghỉ KL", "Giờ làm thiếu",
               "Quy đổi ngày nghỉ", "% Thưởng", "Chi tiết"]
_I_BANG = KHOI_HEADER.index("Bảng chấm công")
KHOI_COT_XUONG_DONG = [_I_BANG, KHOI_HEADER.index("Chi tiết")]


def ten_bang(so: int) -> str:
    """Tên gọi bảng chấm công `so` (1 = NT/TK, 2 = OCP) - ghi ở cột "Bảng chấm công"."""
    return config.CHAM_CONG_TEN_BANG.get(so) or f"Bảng chấm công {so}"


def id_bang(so: int) -> str:
    return (config.GOOGLE_SHEET_ID_CHAM_CONG if so == 1
            else config.GOOGLE_SHEET_ID_CHAM_CONG_2)


def duong_dan(so: int = 1) -> str:
    """Link mở bảng chấm công `so` (nút trên app + dòng Nguồn dưới bảng BC02)."""
    fid = id_bang(so)
    return f"https://docs.google.com/spreadsheets/d/{fid}" if fid else ""


# ----------------------------------------------------------------------
# Tải file từ Drive
# ----------------------------------------------------------------------
def _session() -> AuthorizedSession:
    if not config.SERVICE_ACCOUNT_FILE.exists():
        raise SystemExit(f"Không tìm thấy file service account: {config.SERVICE_ACCOUNT_FILE}")
    creds = Credentials.from_service_account_file(
        str(config.SERVICE_ACCOUNT_FILE), scopes=config.GOOGLE_SCOPES)
    return AuthorizedSession(creds)


def _get(session: AuthorizedSession, url: str, fid: str, **params) -> "requests.Response":
    """GET Drive API, tự thử lại khi 429/5xx; 403/404 -> lỗi rõ nghĩa."""
    import requests

    for delay in _RETRY_DELAYS + (None,):
        try:
            r = session.get(url, params=params, timeout=120)
        except (requests.exceptions.ConnectionError, requests.exceptions.Timeout):
            if delay is None:
                raise
            time.sleep(delay)
            continue
        if r.status_code == 200:
            return r
        if r.status_code in _RETRY_STATUS and delay is not None:
            time.sleep(delay)
            continue
        if r.status_code in (403, 404):
            raise PermissionError(
                f"Drive trả {r.status_code}: file chấm công chưa chia sẻ quyền XEM cho "
                f"email service account, hoặc ID sai ({fid})")
        raise RuntimeError(f"Drive trả {r.status_code}: {r.text[:200]}")
    raise RuntimeError("Drive: hết số lần thử lại")


_file_cache: dict[int, tuple[str, bytes]] = {}


def tai_file(so: int = 1) -> tuple[str, bytes]:
    """(tên file, nội dung .xlsx) của bảng `so`. File Excel trên Drive tải thẳng; Google
    Sheet gốc (bảng OCP) thì xuất ra .xlsx - cùng 1 cách đọc. Cache 1 lần/lần chạy."""
    if so in _file_cache:
        return _file_cache[so]
    fid = id_bang(so)
    if not fid:
        raise LookupError(f"chưa cấu hình ID bảng chấm công {so} (.env / config.py)")
    s = _session()
    base = f"https://www.googleapis.com/drive/v3/files/{fid}"
    meta = _get(s, base, fid, fields="name,mimeType", supportsAllDrives="true").json()
    if meta.get("mimeType") == _GSHEET_MIME:
        data = _get(s, base + "/export", fid, mimeType=_XLSX_MIME).content
    else:
        data = _get(s, base, fid, alt="media", supportsAllDrives="true").content
    _file_cache[so] = (meta.get("name", ""), data)
    return _file_cache[so]


# ----------------------------------------------------------------------
# Đọc tab tháng
# ----------------------------------------------------------------------
def _ma(v) -> str:
    """Chuẩn hóa mã trong ô: bỏ dấu, HOA, bỏ khoảng trắng ("CĐ" -> "CD", " kl " -> "KL")."""
    s = str(v).replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return "".join(s.upper().split())


def _so(v) -> float | None:
    """Ô ghi số giờ (8, 7.5, "7,5") -> float; không phải số -> None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _ngay_cot(v) -> tuple[int, int | None] | None:
    """Ô dòng ngày: datetime -> (ngày, tháng); "01/09" -> (1, 9); ô chỉ ghi số 31 -> (31, None)."""
    if v is None or isinstance(v, bool):
        return None
    if hasattr(v, "day") and hasattr(v, "month"):
        return (v.day, v.month)
    m = _NGAY_RE.match(str(v))
    if m:
        return (int(m.group(1)), int(m.group(2)))
    n = _so(v)
    if n is not None and n == int(n) and 1 <= n <= 31:
        return (int(n), None)
    return None


def _cot_ngay(dong_ngay: list, month: int, year: int, dung_thang: bool) -> dict[int, date]:
    """{chỉ số cột: ngày} từ dòng ngày dưới header. dung_thang=True: chỉ lấy ô đúng tháng;
    False: lấy SỐ NGÀY của mọi ô, bỏ qua tháng (dòng ngày chưa cập nhật). Ô chỉ ghi số
    ngày (vd "31" cuối dòng) chỉ nhận khi đứng ngay sau 1 cột ngày; mỗi ngày 1 cột đầu."""
    out: dict[int, date] = {}
    da_co: set[int] = set()
    for j, v in enumerate(dong_ngay):
        dm = _ngay_cot(v)
        if not dm or dm[0] in da_co:
            continue
        if dm[1] is None and (j - 1) not in out:
            continue
        if dung_thang and dm[1] is not None and dm[1] != month:
            continue
        try:
            out[j] = date(year, month, dm[0])
        except ValueError:
            continue
        da_co.add(dm[0])
    return out


def _thang_nam_tren_header(row: list) -> tuple[int | None, int | None]:
    """Dòng header có "THÁNG" rồi số tháng ở ô bên phải, "NĂM" rồi số năm."""
    thang = nam = None
    for j, v in enumerate(row):
        if v is None:
            continue
        key = _ma(v)
        if key in ("THANG", "NAM"):
            for w in row[j + 1:j + 4]:
                n = _so(w)
                if n is not None:
                    if key == "THANG":
                        thang = int(n)
                    else:
                        nam = int(n)
                    break
    return thang, nam


def _thang_tren_ten_tab(title: str) -> int | None:
    """"BCC.9" / "BCC T9" / "Tháng 9 " -> 9; tab khác ("Phép năm", "NS.2026") -> None."""
    m = _TEN_TAB_RE.match(_ma(title))
    return int(m.group(1)) if m else None


def _chon_tab(wb, month: int, year: int):
    """Worksheet của tháng -> (ws, chỉ số dòng header 0-based, cảnh báo | None).
    Ưu tiên: tên tab VÀ ô THÁNG/NĂM cùng đúng tháng > chỉ ô THÁNG/NĂM đúng (tên tab không
    ghi tháng) > chỉ tên tab đúng (ô THÁNG chưa sửa -> cảnh báo)."""
    theo_header = theo_ten = None
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(min_row=1, max_row=12, values_only=True))
        h = next((i for i, r in enumerate(rows)
                  if any(v is not None and _ma(v) == "HOVATEN" for v in r)), None)
        if h is None:
            continue
        t, n = _thang_nam_tren_header(list(rows[h]))
        header_dung = t is not None and t == month and (n is None or n == year)
        ten_thang = _thang_tren_ten_tab(ws.title)
        if ten_thang == month and (header_dung or t is None):
            return ws, h, None
        if ten_thang is None and header_dung and theo_header is None:
            theo_header = (ws, h, None)
        if ten_thang == month and not header_dung and theo_ten is None:
            theo_ten = (ws, h, f"tab '{ws.title}': ô THÁNG/NĂM trên header ghi "
                               f"{t}/{n if n is not None else '?'} khác tên tab (chưa cập nhật) "
                               f"- tin theo tên tab = tháng {month}/{year}")
    if theo_header:
        return theo_header
    if theo_ten:
        return theo_ten
    raise LookupError(f"file chấm công không có tab tháng {month}/{year} "
                      f"(các tab: {', '.join(ws.title for ws in wb.worksheets)})")


_cache: dict[tuple[int, int, int], dict] = {}


def doc_bang(month: int, year: int, so: int = 1) -> dict:
    """Đọc + tính bảng chấm công `so` của tháng. Trả về:
    {"so", "ten_bang", "file": tên file, "tab": tên tab, "ngay": [ngày có cột] (date),
     "canh_bao": [chuỗi] (ô THÁNG / dòng ngày chưa cập nhật, tên trùng ...),
     "nguoi": [{"ten", "bo_phan", "kl", "thieu_gio", "ngay_quy_doi", "pct", "chi_tiet",
               "ma_la"}] theo thứ tự trên bảng,
     "reg": {tên chuẩn hóa: người - dòng ĐẦU nếu trùng tên}}
    Cache theo (bảng, tháng, năm) trong 1 lần chạy."""
    key = (so, month, year)
    if key in _cache:
        return _cache[key]
    import openpyxl

    ten_file, data = tai_file(so)
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        ws, h, cb_tab = _chon_tab(wb, month, year)
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    canh_bao: list[str] = [cb_tab] if cb_tab else []
    header = list(rows[h])
    i_ten = next(j for j, v in enumerate(header) if v is not None and _ma(v) == "HOVATEN")
    i_bp = next((j for j, v in enumerate(header) if v is not None and _ma(v) == "BOPHAN"), None)

    # Cột ngày = dòng ngay dưới header; chỉ lấy ngày thuộc tháng, mỗi ngày 1 cột đầu tiên
    dong_ngay = list(rows[h + 1]) if h + 1 < len(rows) else []
    cot_ngay = _cot_ngay(dong_ngay, month, year, dung_thang=True)
    if not cot_ngay:
        # Dòng ngày chưa cập nhật (tab "Tháng 9" của OCP vẫn ghi ngày tháng 8):
        # tab đã chọn đúng tháng -> hiểu các cột là ngày 1..31 của tháng này
        cot_ngay = _cot_ngay(dong_ngay, month, year, dung_thang=False)
        if cot_ngay:
            thang_ghi = sorted({dm[1] for v in dong_ngay
                                if (dm := _ngay_cot(v)) and dm[1] is not None})
            canh_bao.append(
                f"tab '{ws.title}': dòng ngày dưới header ghi tháng "
                f"{'/'.join(str(t) for t in thang_ghi) or '?'} (chưa cập nhật) - hiểu các cột "
                f"là ngày {min(cot_ngay.values()).day}..{max(cot_ngay.values()).day} "
                f"của tháng {month:02d}/{year}")
    if not cot_ngay:
        raise LookupError(f"tab '{ws.title}' không có cột ngày của tháng {month}/{year}")

    nguoi: list[dict] = []
    reg: dict[str, dict] = {}
    ten_trung: list[str] = []
    for r in rows[h + 2:]:
        r = list(r)
        stt = _so(r[0]) if r else None
        ten = " ".join(str(r[i_ten]).split()) if i_ten < len(r) and r[i_ten] else ""
        if stt is None or not ten:
            continue                      # dòng trống / tiêu đề khối phụ / "Số lượng nhân sự"
        p = _tinh_nguoi(ten, r, i_bp, cot_ngay)
        nguoi.append(p)
        k = chuan_hoa_ten(ten)
        if k in reg:
            ten_trung.append(ten)         # tên xuất hiện ở 2 khối (vd cả OCP lẫn Tâm An)
        else:
            reg[k] = p
    if ten_trung:
        canh_bao.append("tên xuất hiện 2 lần trên tab (lấy dòng đầu): "
                        + ", ".join(dict.fromkeys(ten_trung)))
    _cache[key] = {"so": so, "ten_bang": ten_bang(so), "file": ten_file, "tab": ws.title,
                   "ngay": sorted(cot_ngay.values()), "canh_bao": canh_bao,
                   "nguoi": nguoi, "reg": reg}
    return _cache[key]


def _tinh_nguoi(ten: str, r: list, i_bp: int | None, cot_ngay: dict[int, date]) -> dict:
    gio_chuan = config.CHAM_CONG_GIO_CHUAN
    kl = 0.0
    thieu_gio = 0.0
    ngay_kl: list[str] = []
    ngay_thieu: list[str] = []
    ma_la: list[str] = []
    for j, d in sorted(cot_ngay.items()):
        v = r[j] if j < len(r) else None
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        n = _so(v)
        if n is not None:
            if 0 <= n < gio_chuan:
                thieu_gio += gio_chuan - n
                ngay_thieu.append(f"{d.strftime('%d/%m')} ({n:g}h)")
            continue
        ma = _ma(v)
        if ma in config.CHAM_CONG_MA_KHONG_LUONG:
            so_ngay = config.CHAM_CONG_MA_KHONG_LUONG[ma]
            kl += so_ngay
            ngay_kl.append(d.strftime("%d/%m") + ("" if so_ngay == 1 else f" ({ma})"))
        elif ma in config.CHAM_CONG_MA_CO_LUONG:
            continue
        else:
            ma_la.append(f"{d.strftime('%d/%m')}: '{str(v).strip()}'")
    ngay_quy_doi = kl + thieu_gio / gio_chuan
    bac = math.floor(ngay_quy_doi / config.CHAM_CONG_NGAY_NGHI_MOI_BAC + 1e-9)
    pct = max(0, 100 - bac * config.CHAM_CONG_TRU_MOI_BAC)
    chi_tiet = []
    if ngay_kl:
        chi_tiet.append("KL: " + ", ".join(ngay_kl))
    if ngay_thieu:
        chi_tiet.append("thiếu giờ: " + ", ".join(ngay_thieu))
    if ma_la:
        chi_tiet.append("mã lạ bỏ qua: " + ", ".join(ma_la))
    return {"ten": ten,
            "bo_phan": " ".join(str(r[i_bp]).split()) if i_bp is not None and i_bp < len(r) and r[i_bp] else "",
            "kl": kl, "thieu_gio": thieu_gio, "ngay_quy_doi": ngay_quy_doi,
            "pct": pct, "chi_tiet": " | ".join(chi_tiet), "ma_la": ma_la}


# ----------------------------------------------------------------------
# Đọc cả 2 bảng
# ----------------------------------------------------------------------
_cac_bang_cache: dict[tuple[int, int], dict] = {}


def doc_cac_bang(month: int, year: int) -> dict:
    """Đọc lần lượt các bảng chấm công của tháng (bảng lỗi KHÔNG chặn bảng kia).
    Trả về {"bang": [kết quả doc_bang của các bảng đọc được, theo thứ tự tìm tên],
            "loi": {số bảng: chuỗi lỗi}}. Cache 1 lần/lần chạy (kể cả lỗi)."""
    key = (month, year)
    if key in _cac_bang_cache:
        return _cac_bang_cache[key]
    kq: dict = {"bang": [], "loi": {}}
    for so in CAC_BANG:
        try:
            kq["bang"].append(doc_bang(month, year, so))
        except Exception as e:                    # noqa: BLE001 - lỗi từng bảng chỉ cảnh báo
            kq["loi"][so] = f"{type(e).__name__}: {e}"
    _cac_bang_cache[key] = kq
    return kq


def _tim(cac_bang: list[dict], ten_pancake: str) -> tuple[dict, dict] | None:
    """Tìm 1 nhân viên Pancake trên các bảng -> (bảng, người) hoặc None.
    Bảng 1 trước; khớp ở nhiều bảng thì lấy tên khớp DÀI nhất (bằng nhau -> bảng đứng trước)."""
    best: tuple[str, dict] | None = None
    for bang in cac_bang:
        ln = khop_ten(bang["reg"], ten_pancake)
        if ln is not None and (best is None or len(ln) > len(best[0])):
            best = (ln, bang)
    if best is None:
        return None
    return best[1], best[1]["reg"][best[0]]


# ----------------------------------------------------------------------
# Áp cho bảng BC02
# ----------------------------------------------------------------------
def pct_theo_roster(roster: list[tuple[str, dict]], month: int, year: int,
                    roster_khoi: list[tuple[str, dict]] | None = None,
                    ) -> tuple[dict[str, str], list[list], list[str], list[str]]:
    """Cho danh sách nhân viên Pancake (uid, {"name","dept"}) của 1 bộ phận:
    - {tên Pancake: "90%"} cho người khớp được trên chấm công (điền cột % Thưởng)
      - tính cho CẢ `roster`; tìm ở bảng 1 trước, không có mới sang bảng 2
    - khối "CHẤM CÔNG" ghi dưới bảng BC02 để soát (mỗi tài khoản Pancake 1 dòng, cột
      theo KHOI_HEADER - có cột "Bảng chấm công" = tên bảng đã lấy) - chỉ gồm
      `roster_khoi` (mặc định = roster; thường là người CÓ đơn chốt / doanh số trong
      tháng, ai cả hai = 0 thì bỏ ra ngoài khối), STT đánh lại từ 1
    - cảnh báo (mã lạ; không khớp - chỉ xét người trong khối)
    - tên (trong khối) KHÔNG có trên bảng chấm công nào - để tô VÀNG dòng đó trên sheet"""
    kq = doc_cac_bang(month, year)
    cac_bang = kq["bang"]
    pct_map: dict[str, str] = {}
    canh_bao: list[str] = []
    da_bao_ma_la: set[tuple[int, str]] = set()
    for _uid, info in roster:
        tim = _tim(cac_bang, info["name"])
        if tim is None:
            continue
        bang, p = tim
        pct_map[info["name"]] = f"{p['pct']}%"
        k = (bang["so"], chuan_hoa_ten(p["ten"]))
        if p["ma_la"] and k not in da_bao_ma_la:
            da_bao_ma_la.add(k)
            canh_bao.append(f"'{p['ten']}' ({bang['ten_bang']}) có mã lạ trên chấm công "
                            "(bỏ qua): " + ", ".join(p["ma_la"]))

    chua_doc = [ten_bang(so) for so in kq["loi"]]
    khong_co = ("(không có trên chấm công"
                + (f"; chưa đọc được {', '.join(chua_doc)}" if chua_doc else "") + ")")
    rows: list[list] = []
    khong_khop: list[str] = []
    for idx, (_uid, info) in enumerate(roster if roster_khoi is None else roster_khoi,
                                       start=1):
        tim = _tim(cac_bang, info["name"])
        if tim is None:
            khong_khop.append(info["name"])
            rows.append([idx, info["name"], khong_co, "", "", "", "", "", "'giữ % cũ", ""])
            continue
        bang, p = tim
        rows.append([idx, info["name"], p["ten"], bang["ten_bang"], p["bo_phan"],
                     round(p["kl"], 2), round(p["thieu_gio"], 2), round(p["ngay_quy_doi"], 3),
                     f"'{p['pct']}%", p["chi_tiet"]])
    if khong_khop:
        canh_bao.append("Không có trên chấm công"
                        + (f" (chưa đọc được {', '.join(chua_doc)})" if chua_doc else
                           f" ({' và '.join(b['ten_bang'] for b in cac_bang)})")
                        + " - giữ % Thưởng cũ / 100%: " + ", ".join(khong_khop))
    return pct_map, rows, canh_bao, khong_khop


def dem_theo_bang(rows: list[list]) -> list[tuple[str, int]]:
    """[(tên bảng, số dòng trong khối lấy từ bảng đó)] - để in log tóm tắt."""
    dem: dict[str, int] = {}
    for r in rows:
        if r[_I_BANG]:
            dem[r[_I_BANG]] = dem.get(r[_I_BANG], 0) + 1
    return [(ten_bang(so), dem[ten_bang(so)]) for so in CAC_BANG if ten_bang(so) in dem]


def build_block(month: int, year: int, group_label: str, rows: list[list],
                so_bo_qua: int = 0) -> list[list]:
    """Khối "CHẤM CÔNG" ghi ngay dưới bảng BC02: tiêu đề, header, dòng nhân viên, quy tắc,
    lưu ý từng bảng, nguồn (mỗi bảng 1 dòng: tab, file, khoảng ngày, link).
    so_bo_qua: số tài khoản không liệt kê (Đơn chốt = 0 và không có doanh số)."""
    kq = doc_cac_bang(month, year)
    ngay = sorted(d for b in kq["bang"] for d in b["ngay"])
    khoang = f"{ngay[0].strftime('%d/%m')} - {ngay[-1].strftime('%d/%m')}" if ngay else ""
    out: list[list] = [
        [f"CHẤM CÔNG THÁNG {month:02d}.{year} - Bộ phận {group_label} -> cột % Thưởng "
         f"(tìm tên ở {ten_bang(1)} trước, không có mới sang {ten_bang(2)}; ngày {khoang}"
         + (f"; chỉ liệt kê người có đơn chốt / doanh số, bỏ qua {so_bo_qua} tài khoản "
            "không có đơn" if so_bo_qua else "") + ")"],
        list(KHOI_HEADER),
    ]
    out += rows
    bac, tru, gio = (config.CHAM_CONG_NGAY_NGHI_MOI_BAC, config.CHAM_CONG_TRU_MOI_BAC,
                     config.CHAM_CONG_GIO_CHUAN)
    out.append([f"Quy tắc: Quy đổi ngày nghỉ = ngày nghỉ KL + giờ làm thiếu / {gio} "
                f"(ngày làm dưới {gio}h thì số giờ thiếu cộng dồn, đủ {gio}h = 1 ngày nghỉ "
                f"không lương). Cứ đủ {bac} ngày -> trừ {tru}% thưởng. "
                "Nghỉ phép (P), nghỉ lễ (NL), nghỉ chế độ (CĐ) có lương -> không trừ. "
                f"Cột \"Bảng chấm công\" = bảng đã lấy số liệu ({ten_bang(1)} = bảng 1, "
                f"{ten_bang(2)} = bảng 2). Không có trên cả 2 bảng -> giữ % đang có trên sheet, "
                "DÒNG TÔ VÀNG (cả trên bảng BC02) - kiểm tra lại tên trên Pancake / chấm công."])
    for b in kq["bang"]:
        for cb in b["canh_bao"]:
            out.append([f"Lưu ý {b['ten_bang']}: {cb}"])
    for so in CAC_BANG:
        b = next((b for b in kq["bang"] if b["so"] == so), None)
        if b:
            n = b["ngay"]
            mo_ta = (f"tab \"{b['tab'].strip()}\" file \"{b['file']}\", ngày "
                     f"{n[0].strftime('%d/%m')} - {n[-1].strftime('%d/%m')}" if n else
                     f"tab \"{b['tab'].strip()}\" file \"{b['file']}\"")
        elif so in kq["loi"]:
            mo_ta = f"KHÔNG ĐỌC ĐƯỢC ({kq['loi'][so][:120]})"
        else:
            continue
        out.append([f"Nguồn {ten_bang(so)}: {mo_ta}"
                    + (f" - {duong_dan(so)}" if duong_dan(so) else "")])
    return out
